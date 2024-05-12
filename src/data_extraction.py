import json
import csv
from googleapiclient.discovery import build
from PyPDF2 import PdfReader
import openpyxl
from docx import Document

def extract_from_gdoc(api_service, file_id):
    # Use Google Drive API to download the file as text
    pass

def extract_from_gsheet(api_service, file_id):
    # Use Google Drive API to download the file as CSV
    pass

def extract_from_json(file_path):
    with open(file_path, 'r') as file:
        return json.load(file)

def extract_from_csv(file_path):
    with open(file_path, newline='') as file:
        return list(csv.DictReader(file))

def extract_from_pdf(file_path):
    reader = PdfReader(file_path)
    text = [page.extract_text() for page in reader.pages]
    return text

def extract_from_xlsx(file_path):
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    return data

def extract_from_docx(file_path):
    doc = Document(file_path)
    return [para.text for para in doc.paragraphs]

def extract_from_txt(file_path):
    with open(file_path, 'r') as file:
        return file.read()

